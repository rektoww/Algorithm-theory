from datetime import datetime
from pathlib import Path

from docx import Document

from .content import Exhibit
from .analytics import get_visitor_statistics, rank_content


def export_report(contents, visitors, output_dir=None):
    # Папка для созданных отчётов
    if output_dir is None:
        reports_dir = Path(__file__).resolve().parent.parent / "reports"
    else:
        reports_dir = Path(output_dir)

    reports_dir.mkdir(parents=True, exist_ok=True)

    # Уникальное имя документа
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_path = reports_dir / f"museum_report_{timestamp}.docx"

    document = Document()

    document.add_heading("Виртуальный музей", 0)
    document.add_paragraph(
        f"Отчёт сформирован: {datetime.now():%d.%m.%Y %H:%M}"
    )

    # Общая информация
    document.add_heading("Общая статистика", level=1)
    document.add_paragraph(f"Количество материалов: {len(contents)}")
    document.add_paragraph(f"Количество посетителей: {len(visitors)}")

    # Рейтинг материалов
    document.add_heading("Популярность материалов", level=1)

    ranked = rank_content(contents.values())

    if ranked:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"

        headers = [
            "ID", "Тип", "Название", "Просмотры", "Популярность"
        ]

        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header

        for content in ranked:
            row = table.add_row().cells

            row[0].text = str(content.content_id)
            row[1].text = (
                "Экспонат" if isinstance(content, Exhibit)
                else "Статья"
            )
            row[2].text = content.title
            row[3].text = str(content.views)
            row[4].text = str(content.calculate_popularity())

    else:
        document.add_paragraph("В музее пока нет материалов.")

    # Статистика посетителей
    document.add_heading("Статистика посетителей", level=1)

    labels = {
        "exhibit_views": "Просмотров экспонатов",
        "unique_exhibits": "Уникальных экспонатов",
        "article_reads": "Прочитано статей",
        "unique_articles": "Уникальных статей",
        "likes_given": "Поставлено лайков",
        "favorites_given": "Добавлено в избранное",
        "favorites_added": "Добавлено в избранное",
        "bookmarks_added": "Добавлено закладок",
        "total_actions": "Всего действий"
    }

    if not visitors:
        document.add_paragraph("Посетители отсутствуют.")

    for visitor in visitors.values():
        document.add_heading(
            f"{visitor.name} (ID: {visitor.visitor_id})",
            level=2
        )

        statistics = get_visitor_statistics(visitor)

        for key, value in statistics.items():
            if key in ("visitor_id", "name"):
                continue

            label = labels.get(key, key)
            document.add_paragraph(f"{label}: {value}")

    document.save(file_path)

    return file_path