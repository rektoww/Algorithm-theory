from docx import Document

from virtual_museum.content import Exhibit, Article
from virtual_museum.visitor import Visitor
from virtual_museum.report import export_report


def test_export_report_with_data(tmp_path):
    exhibit = Exhibit(1, "Мона Лиза", "Живопись")
    article = Article(2, "История музея", "Иван Иванов")
    visitor = Visitor(1, "Максим")

    visitor.view_exhibit(exhibit)
    visitor.like_content(exhibit)
    visitor.favorite_exhibit(exhibit)

    contents = {
        1: exhibit,
        2: article
    }

    visitors = {
        1: visitor
    }

    file_path = export_report(
        contents,
        visitors,
        output_dir=tmp_path
    )

    assert file_path.exists()
    assert file_path.suffix == ".docx"

    document = Document(file_path)

    paragraphs = [p.text for p in document.paragraphs]

    assert "Количество материалов: 2" in paragraphs
    assert "Количество посетителей: 1" in paragraphs
    assert "Всего действий: 3" in paragraphs

    assert len(document.tables) == 1

    table = document.tables[0]

    assert len(table.rows) == 3
    assert table.cell(1, 2).text == "Мона Лиза"
    assert table.cell(1, 4).text == "6"
    assert table.cell(2, 2).text == "История музея"


def test_export_empty_report(tmp_path):
    file_path = export_report(
        {},
        {},
        output_dir=tmp_path
    )

    assert file_path.exists()

    document = Document(file_path)
    paragraphs = [p.text for p in document.paragraphs]

    assert "В музее пока нет материалов." in paragraphs
    assert "Посетители отсутствуют." in paragraphs