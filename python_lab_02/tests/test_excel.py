from openpyxl import load_workbook

from virtual_museum.content import Exhibit, Article
from virtual_museum.visitor import Visitor
from virtual_museum.report import export_excel


def test_export_excel(tmp_path):
    exhibit = Exhibit(1, "Мона Лиза", "Живопись")
    article = Article(2, "История музея", "Иван Иванов")
    visitor = Visitor(1, "Максим")

    visitor.view_exhibit(exhibit)
    visitor.like_content(exhibit)
    visitor.favorite_exhibit(exhibit)
    visitor.read_article(article)
    visitor.bookmark_article(article)

    contents = {1: exhibit, 2: article}
    visitors = {1: visitor}

    file_path = export_excel(
        contents,
        visitors,
        output_dir=tmp_path
    )

    assert file_path.exists()
    assert file_path.suffix == ".xlsx"

    workbook = load_workbook(file_path)

    assert workbook.sheetnames == [
        "Сводка",
        "Популярность",
        "Посетители"
    ]

    summary = workbook["Сводка"]
    assert summary["B2"].value == 2
    assert summary["B3"].value == 1

    materials = workbook["Популярность"]
    assert materials["C2"].value == "Мона Лиза"
    assert materials["G2"].value == 6
    assert materials["C3"].value == "История музея"
    assert materials["G3"].value == 4

    visitors_sheet = workbook["Посетители"]
    assert visitors_sheet["B2"].value == "Максим"
    assert visitors_sheet["J2"].value == 5

    workbook.close()


def test_export_empty_excel(tmp_path):
    file_path = export_excel({}, {}, output_dir=tmp_path)

    assert file_path.exists()

    workbook = load_workbook(file_path)

    summary = workbook["Сводка"]
    assert summary["B2"].value == 0
    assert summary["B3"].value == 0

    workbook.close()